#!/usr/bin/env python
# -*- coding: UTF-8 -*-
import os
import pandas as pd
from typing import Dict, Any
from openpyxl import load_workbook
from app.config.setting import test_param
from app.config.logger_setup import logger


class ExcelUtils:
    def __init__(self):
        self.test_enable = test_param.test_enable
        self.output_file = test_param.test_output_file

    def _ensure_directory_exists(self):
        directory = os.path.dirname(self.output_file)
        if not os.path.exists(directory):
            os.makedirs(directory)

    def _ensure_file_exists(self, sheet_name: str):
        self._ensure_directory_exists()
        if not os.path.exists(self.output_file):
            with pd.ExcelWriter(self.output_file, engine='openpyxl') as writer:
                pd.DataFrame().to_excel(writer, sheet_name=sheet_name, index=False)
        else:
            workbook = load_workbook(self.output_file)
            if sheet_name not in workbook.sheetnames:
                with pd.ExcelWriter(self.output_file, engine='openpyxl', mode='a') as writer:
                    pd.DataFrame().to_excel(writer, sheet_name=sheet_name, index=False)

    def export_result(self, result: Dict[str, Any], sheet_name: str):
        if not self.test_enable:
            return
        try:
            self._ensure_file_exists(sheet_name)
            try:
                existing_df = pd.read_excel(self.output_file, sheet_name=sheet_name)
            except Exception:
                existing_df = pd.DataFrame()
            new_df = pd.DataFrame([result])
            combined_df = pd.concat([existing_df, new_df], ignore_index=True)
            with pd.ExcelWriter(self.output_file, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
                combined_df.to_excel(writer, sheet_name=sheet_name, index=False)
        except Exception as e:
            logger.error(f"Failed to export result to Excel: {e}")
